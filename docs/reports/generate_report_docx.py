import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles & Colors
    # Primary: Deep Navy (#0B1E2D), Accent: Vivid Blue (#2563EB), Warm Red/Orange (#FF6B2C), Slate: (#475569)
    COLOR_PRIMARY = RGBColor(11, 30, 45)      # #0B1E2D
    COLOR_ACCENT = RGBColor(37, 99, 235)      # #2563EB
    COLOR_MUTED = RGBColor(71, 85, 105)       # #475569
    COLOR_DARK = RGBColor(15, 23, 42)         # #0F172A

    # Set normal style font
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Segoe UI'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = COLOR_DARK
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # TITLE & HEADER
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("DEADSTOCK AI PLATFORM")
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Comprehensive Project Report & Solution Architecture Specification")
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = COLOR_ACCENT
    sub_p.paragraph_format.space_after = Pt(20)

    # Meta Info Box Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Project Name:", "DeadStock (AI-Powered B2B Circular Liquidation Platform)"),
        ("Version & Status:", "Version 1.0.0 (Full-Stack Next.js 16 + Laravel 12 Architecture)"),
        ("Domain / Industry:", "Industrial Supply Chain, B2B Inventory Risk & Circular Economy"),
        ("Prepared For:", "Executive Leadership, Stakeholders & Engineering Teams"),
    ]

    for row_idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[row_idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.0)
        cell_val.width = Inches(4.5)
        
        p_l = cell_lbl.paragraphs[0]
        r_l = p_l.add_run(label)
        r_l.font.bold = True
        r_l.font.size = Pt(9.5)
        r_l.font.color.rgb = COLOR_PRIMARY

        p_v = cell_val.paragraphs[0]
        r_v = p_v.add_run(val)
        r_v.font.size = Pt(9.5)
        r_v.font.color.rgb = COLOR_MUTED

        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "F8FAFC")
        set_cell_margins(cell_lbl, 80, 80, 100, 100)
        set_cell_margins(cell_val, 80, 80, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    # Helper function for adding styled headings
    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = COLOR_PRIMARY
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_ACCENT
        return h

    def add_bullet(bold_prefix, text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        r_pre = p.add_run(bold_prefix)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_PRIMARY
        r_txt = p.add_run(f" {text}")
        r_txt.font.color.rgb = COLOR_DARK
        return p

    # -------------------------------------------------------------
    # 1. EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    add_heading_1("1. Executive Summary")
    p = doc.add_paragraph(
        "DeadStock is a next-generation, AI-driven B2B circular inventory monetization and risk mitigation platform. "
        "Engineered specifically for industrial manufacturers, distributors, component assemblers, and supply chain enterprises, "
        "DeadStock prevents dead inventory before it crystallizes into complete balance sheet write-offs."
    )
    p = doc.add_paragraph(
        "By fusing predictive degradation algorithms, 5-factor buyer matchmaking, 8-month multi-channel recovery simulations, "
        "and integrated freight crating logistics, the platform enables enterprises to unlock up to 70%–85% of dormant working capital "
        "while diverting thousands of tons of industrial surplus from landfills."
    )

    # -------------------------------------------------------------
    # 2. PROBLEM STATEMENT & INDUSTRY CHALLENGES
    # -------------------------------------------------------------
    add_heading_1("2. Problem Statement & The Reality of Dead Inventory")
    doc.add_paragraph(
        "In modern manufacturing and distribution, dead stock is rarely an instantaneous event. It is a slow, compound financial bleed. "
        "Industrial supply chains face critical bottlenecks:"
    )

    add_bullet("Holding Cost Hemorrhage:", "Inventory carries an average 18% to 24% annual carrying cost bleed (warehousing footprint, capital interest, insurance, labor, and depreciation).")
    add_bullet("Order Cancellations & Specification Shifts:", "Sudden client PO cancellations leave custom-machined or bulk batches with zero active downstream run-rate, freezing operating liquidity.")
    add_bullet("Forecasting Inaccuracies & MOQs:", "Procurement teams are forced to purchase minimum order quantities based on optimistic projections, causing 4x to 10x supply overhangs.")
    add_bullet("Component Obsolescence:", "Rapid firmware changes, revised engineering tolerances, or statutory environmental changes render inventory obsolete before traditional sales channels can react.")
    add_bullet("Fragmented Scrap Channels:", "Without structured digital marketplaces, companies liquidate surplus through local scrap dealers for pennies on the dollar (5%–10% salvage value).")

    # Problem Matrix Table
    table_prob = doc.add_table(rows=5, cols=4)
    table_prob.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    headers = ["Problem Vector", "Root Cause", "Financial Impact", "DeadStock Countermeasure"]
    for i, h in enumerate(headers):
        cell = table_prob.rows[0].cells[i]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "0B1E2D")
        set_cell_margins(cell, 100, 100, 100, 100)

    prob_rows = [
        ("Velocity Stagnation", "Gradual drop in demand", "18-24% annual holding bleed", "Early AI risk velocity alerts at 60-day mark"),
        ("Cancelled Orders", "Downstream buyer pullback", "100% locked working capital", "Instant automated B2B buyer matchmaking"),
        ("MOQ Overhang", "Bulk procurement rules", "Severe balance sheet liability", "Inter-depot transfer & secondary marketplace"),
        ("Component EOL", "Spec & firmware revisions", "Total scrap write-off (0%)", "Multi-channel recovery simulation & recycling")
    ]

    for r_idx, row_data in enumerate(prob_rows, start=1):
        for c_idx, val in enumerate(row_data):
            cell = table_prob.rows[r_idx].cells[c_idx]
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 80, 80, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 3. PROPOSED SOLUTION & PLATFORM ARCHITECTURE
    # -------------------------------------------------------------
    add_heading_1("3. The DeadStock Solution Architecture")
    doc.add_paragraph(
        "DeadStock operates on a dual-engine paradigm: Upstream Predictive Risk Triage combined with Downstream Circular Liquidation Execution. "
        "The system continuously audits enterprise inventory velocity and maps surplus to active manufacturing demand."
    )

    add_heading_2("Core Subsystems & Functional Modules")
    add_bullet("1. Inventory Ingestion & Health Engine:", "Connects to enterprise ERP systems or ingests bulk CSV datasets, automatically evaluating age, unit costs, monthly demand, and degradation risk.")
    add_bullet("2. Multi-Channel Recovery Calculator:", "Simulates financial return across 5 recovery strategies: Direct B2B Liquidation, Inter-depot Transfer, Vendor Restock Claim, Discounted Flash Clearance, and Certified Circular Recycling.")
    add_bullet("3. 5-Factor Algorithmic Matchmaker:", "Ranks potential buyer enterprises using composite vectors: Product Compatibility, Quantity Tolerance, Industry Domain, Geographic Logistics Proximity, and Price Alignment.")
    add_bullet("4. Live Circular Marketplace:", "A secure, verified B2B trading floor for surplus industrial components, polymer connectors, PCBs, and packaging materials.")
    add_bullet("5. Pack-and-Ship Freight Logistics:", "Automated crating guides, freight dispatch coordination, shipment tracking, and escrow settlement.")
    add_bullet("6. ESG & Circular Sustainability Ledger:", "Calculates avoided landfill tonnage, recirculated component counts, and Scope 3 greenhouse gas offsets.")
    add_bullet("7. DeadStock AI Intelligence Assistant:", "Natural language conversational agent that parses user intent, prioritizes emergency liquidation, and generates instant executive summaries.")

    # -------------------------------------------------------------
    # 4. KEY FEATURES & CAPABILITIES BREAKDOWN
    # -------------------------------------------------------------
    add_heading_1("4. In-Depth Feature Specification")

    add_heading_2("4.1. Predictive Risk Scoring Model")
    doc.add_paragraph(
        "Each SKU is evaluated continuously and assigned an AI Risk Score (0% to 100%) and a categorical risk tier:"
    )
    add_bullet("CRITICAL (80% - 100%):", "Stock level exceeds 3.5x annual run-rate with velocity drop > 50%. Immediate liquidation protocol triggered.")
    add_bullet("HIGH (60% - 79%):", "Holding costs exceed 15% of margin. Secondary buyer matching recommended.")
    add_bullet("MODERATE (40% - 59%):", "Turnover rate decelerating. Recommended action: Monitor or apply localized promotional discount.")
    add_bullet("LOW (0% - 39%):", "Healthy operating stock within standard procurement cycle turnover.")

    add_heading_2("4.2. Recovery Yield Comparison Matrix")
    
    table_rec = doc.add_table(rows=6, cols=4)
    table_rec.alignment = WD_TABLE_ALIGNMENT.CENTER
    rec_headers = ["Recovery Strategy", "Capital Recaptured", "Execution Time", "Primary Application"]
    for i, h in enumerate(rec_headers):
        cell = table_rec.rows[0].cells[i]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "0B1E2D")
        set_cell_margins(cell, 100, 100, 100, 100)

    rec_data = [
        ("Direct B2B Buyer Sale", "70% – 85% of Cost", "3 – 7 Days", "High-spec components matching active buyers"),
        ("Inter-Depot Redistribution", "90% – 95% of Value", "2 – 4 Days", "Stock deficit in sister warehouses"),
        ("Vendor Return / Buyback", "75% – 85% (less fee)", "10 – 14 Days", "Active vendor restock agreements"),
        ("Marketplace Flash Clearance", "50% – 65% of Cost", "5 – 10 Days", "Standardized commodity hardware"),
        ("Certified Circular Recycling", "25% – 40% (Material)", "1 – 3 Days", "Deprecated plastics, metals, and scrap")
    ]

    for r_idx, row_data in enumerate(rec_data, start=1):
        for c_idx, val in enumerate(row_data):
            cell = table_rec.rows[r_idx].cells[c_idx]
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 80, 80, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 5. TECHNOLOGY STACK & API ARCHITECTURE
    # -------------------------------------------------------------
    add_heading_1("5. Technology Stack & Technical Implementation")

    doc.add_paragraph(
        "DeadStock is built with a resilient, modern, high-performance decoupled architecture:"
    )

    add_bullet("Frontend Framework:", "Next.js 16 (App Router + Turbopack) leveraging React 19, TypeScript, and Server/Client Components.")
    add_bullet("Styling & UX:", "Tailwind CSS v4, Framer Motion for kinetic physics, GSAP, and Lucide React icons.")
    add_bullet("Backend API Layer:", "Laravel 12 (PHP 8.3) with RESTful API v1 architecture, Eloquent ORM, and database-agnostic drivers (PostgreSQL / SQLite).")
    add_bullet("Security & Auth:", "Laravel Sanctum token-based authentication with role-based access control (Admin, Seller, Buyer, Auditor).")
    add_bullet("Analytical / ML Bridge:", "Python / Scikit-Learn analytical scripts for demand forecasting and multi-dimensional regression modeling.")

    add_heading_2("Core API Endpoints Overview")
    add_bullet("Authentication:", "POST /api/v1/auth/register, POST /api/v1/auth/login, GET /api/v1/auth/me")
    add_bullet("Inventory & Risk:", "GET /api/v1/inventory (with filters), GET /api/v1/inventory/{id}, POST /api/v1/inventory")
    add_bullet("Buyer Matchmaking:", "GET /api/v1/matches, GET /api/v1/matches/{sku}")
    add_bullet("Marketplace Operations:", "GET /api/v1/marketplace/listings, POST /api/v1/marketplace/listings")
    add_bullet("Analytics & ESG:", "GET /api/v1/dashboard/summary, GET /api/v1/sustainability/metrics")
    add_bullet("AI Intelligence:", "POST /api/v1/ai/ask (Intent-driven conversational NLP engine)")

    # -------------------------------------------------------------
    # 6. BUSINESS MODEL & PRICING TIERS
    # -------------------------------------------------------------
    add_heading_1("6. Business Model & Commercial Strategy")
    doc.add_paragraph(
        "DeadStock deploys a diversified SaaS + Transaction fee monetization model designed to lower adoption barriers while capturing value upon successful liquidations."
    )

    table_price = doc.add_table(rows=4, cols=4)
    table_price.alignment = WD_TABLE_ALIGNMENT.CENTER
    p_headers = ["Plan Tier", "Monthly / Annual Price", "Included Listings & Capacity", "Target Customer"]
    for i, h in enumerate(p_headers):
        cell = table_price.rows[0].cells[i]
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "0B1E2D")
        set_cell_margins(cell, 100, 100, 100, 100)

    price_data = [
        ("Basic Tier", "₹0 / Free", "Up to 50 active listings, CSV import, basic analytics", "Small businesses & single-facility suppliers"),
        ("Pro Plan", "₹4,999/mo (₹3,999/mo yearly)", "Unlimited listings, AI buyer matching, 50 shipments, 24/7 support", "Mid-sized manufacturers & active shippers"),
        ("Enterprise", "Custom Quote", "Unlimited shipments, ERP/SAP integration, white-label, custom logistics", "Multi-national industrial enterprises & 3PLs")
    ]

    for r_idx, row_data in enumerate(price_data, start=1):
        for c_idx, val in enumerate(row_data):
            cell = table_price.rows[r_idx].cells[c_idx]
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            set_cell_background(cell, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(cell, 80, 80, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 7. STRATEGIC ADVANTAGES & ROADMAP
    # -------------------------------------------------------------
    add_heading_1("7. Strategic Advantages & Future Roadmap")

    add_heading_2("Key Competitive Differentiators")
    add_bullet("Proactive vs Reactive:", "Identifies risk in months 3-6 rather than writing off at month 18.")
    add_bullet("High-Precision Matching:", "5-factor compatibility removes guesswork for industrial procurement teams.")
    add_bullet("Auditable ESG Impact:", "Provides direct Scope 3 carbon offset figures and landfill diversion certs.")

    add_heading_2("Strategic Roadmap Milestones")
    add_bullet("Phase 1 (Completed):", "Core Next.js frontend, Laravel 12 API, AI risk model, and marketplace.")
    add_bullet("Phase 2 (Upcoming):", "Direct bi-directional webhooks for SAP S/4HANA, Oracle NetSuite, and Microsoft Dynamics 365.")
    add_bullet("Phase 3 (Expansion):", "Cross-border customs and GST-compliant e-invoicing for international freight liquidation.")
    add_bullet("Phase 4 (Advanced):", "IoT-verified smart warehouse delivery contracts with instant escrow settlement.")

    # -------------------------------------------------------------
    # CONCLUSION
    # -------------------------------------------------------------
    add_heading_1("8. Conclusion")
    doc.add_paragraph(
        "DeadStock represents a structural upgrade to industrial supply chain management. "
        "By replacing opaque liquidation auctions with predictive risk intelligence and direct B2B buyer matchmaking, "
        "DeadStock safeguards profit margins, optimizes warehouse real estate, and fosters a zero-waste industrial economy."
    )

    # Save document
    output_path = r"c:\Users\Dhruv\OneDrive\Desktop\DeadStock\DeadStock_Project_Report.docx"
    doc.save(output_path)
    print(f"Report saved successfully at: {output_path}")

if __name__ == "__main__":
    create_report()
